from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.database import SessionLocal
from app.main import app
from app.matching import normalize_text
from app.models import HistoryQuote


def workbook_bytes(rows: list[list] | None = None) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "询价单"
    sheet.append(["序号", "产品名称", "参数", "型号", "单位", "数量"])
    for row in (rows or [[1, "电子天平", "100g，0.001g，带防风罩", "", "台", 2]]):
        sheet.append(row)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def login(client: TestClient, username: str = "admin", password: str = "admin123"):
    response = client.post("/api/auth/login", json={"username": username, "password": password})
    assert response.status_code == 200, response.text


def ordinary_customer_id(client: TestClient) -> int:
    customers = client.get("/api/customers").json()
    return next(item for item in customers if item["customer_type"] == "ordinary")["id"]


def create_job(client: TestClient, customer_id: int | None, content: bytes | None = None) -> str:
    data = {"requested_option_count": 3}
    if customer_id is not None:
        data["customer_id"] = customer_id
    response = client.post(
        "/api/quote-jobs",
        data=data,
        files={"file": ("询价.xlsx", content or workbook_bytes(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 202, response.text
    return response.json()["id"]


def first_line_detail(client: TestClient, job_id: str) -> dict:
    lines = client.get(f"/api/quote-jobs/{job_id}/lines").json()
    assert lines["items"], "job produced no lines"
    return client.get(f"/api/quote-lines/{lines['items'][0]['id']}").json()


def add_history(record_id: str, name: str, price: float, spec: str = "100g", model: str = "T-1",
                manufacturer: str = "重复测试厂", unit: str = "台"):
    db = SessionLocal()
    try:
        db.add(
            HistoryQuote(
                id=record_id,
                source_file="手工录入",
                name=name,
                normalized_name=normalize_text(name),
                spec=spec,
                normalized_spec=normalize_text(spec),
                model=model,
                normalized_model=normalize_text(model),
                manufacturer=manufacturer,
                unit=unit,
                normalized_unit=normalize_text(unit),
                price=price,
            )
        )
        db.commit()
    finally:
        db.close()


def test_candidate_dedup_keeps_single_identical_option():
    add_history("manual:dup-a", "去重专用天平X", 500.0)
    add_history("manual:dup-b", "去重专用天平X", 500.0)
    with TestClient(app) as client:
        login(client)
        content = workbook_bytes([[1, "去重专用天平X", "100g", "T-1", "台", 1]])
        job_id = create_job(client, ordinary_customer_id(client), content)
        detail = first_line_detail(client, job_id)
        identical = [
            item for item in detail["options"]
            if item["record"]["id"] in ("manual:dup-a", "manual:dup-b")
        ]
        assert len(identical) == 1
        ranks = [item["rank"] for item in detail["options"]]
        assert ranks == sorted(ranks)


def test_patch_allows_same_manufacturer_different_price():
    add_history("manual:same-maker-a", "同厂不同价天平Y", 500.0)
    add_history("manual:same-maker-b", "同厂不同价天平Y", 700.0)
    with TestClient(app) as client:
        login(client)
        content = workbook_bytes([[1, "同厂不同价天平Y", "100g", "T-1", "台", 1]])
        job_id = create_job(client, ordinary_customer_id(client), content)
        detail = first_line_detail(client, job_id)
        pair = [
            item for item in detail["options"]
            if item["record"]["id"] in ("manual:same-maker-a", "manual:same-maker-b")
        ]
        assert len(pair) == 2
        response = client.patch(
            f"/api/quote-lines/{detail['id']}",
            json={"selected_option_ids": [item["id"] for item in pair], "final_prices": {}, "manual_note": ""},
        )
        assert response.status_code == 200, response.text
        assert [item["selected"] for item in response.json()["options"] if item["id"] in {p["id"] for p in pair}] == [True, True]


def test_patch_rejects_fully_identical_selection():
    with TestClient(app) as client:
        login(client)
        job_id = create_job(client, ordinary_customer_id(client))
        detail = first_line_detail(client, job_id)
        # two manual options with identical manufacturer+price+spec+model
        ids = []
        for _ in range(2):
            created = client.post(
                f"/api/quote-lines/{detail['id']}/options",
                json={"manufacturer": "手工同厂", "brand": "手工牌", "spec": "100g", "model": "M1",
                      "unit": "台", "price": 123.0},
            )
            assert created.status_code == 201, created.text
            ids.append(created.json()["id"])
        response = client.patch(
            f"/api/quote-lines/{detail['id']}",
            json={"selected_option_ids": ids, "final_prices": {}, "manual_note": ""},
        )
        assert response.status_code == 400
        assert "不能重复选择完全相同的方案" in response.json()["detail"]
        # changing one price makes the pair selectable
        ok = client.patch(
            f"/api/quote-lines/{detail['id']}",
            json={"selected_option_ids": ids, "final_prices": {str(ids[1]): 130.0}, "manual_note": ""},
        )
        assert ok.status_code == 200, ok.text


def test_manual_option_lifecycle_and_export():
    with TestClient(app) as client:
        login(client)
        job_id = create_job(client, ordinary_customer_id(client))
        detail = first_line_detail(client, job_id)
        line_id = detail["id"]

        # validation
        bad_price = client.post(f"/api/quote-lines/{line_id}/options", json={"manufacturer": "x", "price": 0})
        assert bad_price.status_code == 422
        no_maker = client.post(f"/api/quote-lines/{line_id}/options", json={"price": 10})
        assert no_maker.status_code == 400

        created = client.post(
            f"/api/quote-lines/{line_id}/options",
            json={"manufacturer": "手工制造厂", "brand": "手工牌", "model": "SG-1",
                  "spec": "定制参数", "unit": "台", "price": 66.6},
        )
        assert created.status_code == 201, created.text
        option = created.json()
        assert option["confidence"] == "manual"
        assert option["selected"] is True
        assert option["final_price"] == 66.6
        assert option["record"]["manufacturer"] == "手工制造厂"
        assert option["record"]["source_file"] == "手工方案"
        assert option["record"]["quote_date"] == ""

        # history-linked options cannot be deleted
        history_option = next(item for item in detail["options"])
        rejected = client.delete(f"/api/quote-options/{history_option['id']}")
        assert rejected.status_code == 400

        # confirm + export with the manual option selected must not crash
        patch = client.patch(
            f"/api/quote-lines/{line_id}",
            json={"selected_option_ids": [option["id"]], "final_prices": {}, "manual_note": ""},
        )
        assert patch.status_code == 200
        confirm = client.post(f"/api/quote-lines/{line_id}/confirm", json={"override_reason": ""})
        assert confirm.status_code == 200, confirm.text
        internal = client.post(f"/api/quote-jobs/{job_id}/export/internal")
        assert internal.status_code == 200
        workbook = load_workbook(BytesIO(internal.content))
        joined = " ".join(str(cell.value or "") for row in workbook["内部方案明细"] for cell in row)
        assert "手工制造厂" in joined
        assert "手工方案" in joined

        deleted = client.delete(f"/api/quote-options/{option['id']}")
        assert deleted.status_code == 200
        detail_after = client.get(f"/api/quote-lines/{line_id}").json()
        assert all(item["id"] != option["id"] for item in detail_after["options"])


def test_confirm_with_blocking_warning_no_longer_requires_reason():
    with TestClient(app) as client:
        login(client)
        created = client.post(
            "/api/customers",
            json={"name": "VIP覆盖原因测试", "customer_type": "vip",
                  "discount_percent": 8, "minimum_margin_percent": 15},
        )
        assert created.status_code == 200
        job_id = create_job(client, created.json()["id"])
        detail = first_line_detail(client, job_id)
        assert any(
            str(warning).startswith("BLOCK:")
            for option in detail["options"] if option["selected"]
            for warning in option["warnings"]
        )
        confirm = client.post(f"/api/quote-lines/{detail['id']}/confirm", json={"override_reason": ""})
        assert confirm.status_code == 200, confirm.text
        assert confirm.json()["confirmed"] is True


def test_job_without_customer():
    with TestClient(app) as client:
        login(client)
        job_id = create_job(client, None)
        job = client.get(f"/api/quote-jobs/{job_id}").json()
        assert job["customer"] is None
        assert job["status"] == "review"
        assert job["total_lines"] == 1
        lines = client.get(f"/api/quote-jobs/{job_id}/lines").json()
        assert lines["total"] == 1


def test_customers_list_created_at_desc():
    with TestClient(app) as client:
        login(client)
        first = client.post("/api/customers", json={"name": "排序客户甲"}).json()
        second = client.post("/api/customers", json={"name": "排序客户乙"}).json()
        customers = client.get("/api/customers").json()
        assert all("created_at" in item for item in customers)
        ids = [item["id"] for item in customers]
        assert ids.index(second["id"]) < ids.index(first["id"])
        timestamps = [item["created_at"] for item in customers]
        assert timestamps == sorted(timestamps, reverse=True)


def history_import_workbook() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "历史"
    sheet.append(["序号", "产品名称", "参数", "单位", "制造商", "含税单价"])
    sheet.append([1, "导入测试烧杯", "500ml", "个", "导入厂A", 12.5])
    sheet.append([2, "导入测试试管", "15ml", "支", "导入厂B", 3.0])
    sheet.append([3, "导入测试无价行", "无价格", "个", "导入厂A", None])
    sheet.append(["", "合计", "", "", "", None])
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def test_import_history_counts_and_dedup():
    with TestClient(app) as client:
        login(client)
        payload = {"file": ("历史库.xlsx", history_import_workbook(),
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        first = client.post("/api/governance/import-history", files=payload)
        assert first.status_code == 200, first.text
        assert first.json() == {"inserted": 2, "skipped_duplicates": 0, "skipped_invalid": 1}

        found = client.get("/api/history/search", params={"q": "导入测试烧杯"}).json()
        assert found and found[0]["price"] == 11.36

        # importing the same file again yields only duplicates
        second = client.post(
            "/api/governance/import-history",
            files={"file": ("历史库.xlsx", history_import_workbook(),
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert second.json()["inserted"] == 0
        assert second.json()["skipped_duplicates"] == 2
        assert second.json()["skipped_invalid"] == 1

        bad = client.post("/api/governance/import-history", files={"file": ("a.txt", b"x", "text/plain")})
        assert bad.status_code == 400


def test_import_history_requires_admin():
    with TestClient(app) as client:
        login(client, "quote", "quote123")
        response = client.post(
            "/api/governance/import-history",
            files={"file": ("历史库.xlsx", history_import_workbook(),
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert response.status_code == 403


def test_import_history_fuzzy_price_header():
    """价格列叫 ``赛特尔单价`` 这类非精确别名时也能导入。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "报价"
    sheet.append(["序号", "器材名称", "规格、品名、教学性能要求", "单位", "数量", "赛特尔单价"])
    sheet.append([1, "模糊价格烧杯", "500ml", "个", 5, 12.5])
    sheet.append([2, "模糊价格试管", "15ml", "支", 10, 3.0])
    stream = BytesIO()
    workbook.save(stream)

    with TestClient(app) as client:
        login(client)
        response = client.post(
            "/api/governance/import-history",
            files={"file": ("高中理化生报价.xlsx", stream.getvalue(),
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert response.status_code == 200, response.text
        assert response.json()["inserted"] == 2
        found = client.get("/api/history/search", params={"q": "模糊价格烧杯"}).json()
        assert found and found[0]["price"] == 12.5


def test_import_history_without_price_column_returns_400():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "清单"
    sheet.append(["序号", "采购品目", "参数", "数量", "单位"])
    sheet.append([1, "无价格烧杯", "500ml", 5, "个"])
    stream = BytesIO()
    workbook.save(stream)

    with TestClient(app) as client:
        login(client)
        response = client.post(
            "/api/governance/import-history",
            files={"file": ("无价格清单.xlsx", stream.getvalue(),
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert response.status_code == 400
        assert "未找到价格列" in response.json()["detail"]

        # 真实的无价格列招标清单同样报 400 而不是静默成功
        real_file = Path(__file__).resolve().parents[3] / "测试数据" / "海南发改委14包.xlsx"
        if real_file.exists():
            response = client.post(
                "/api/governance/import-history",
                files={"file": (real_file.name, real_file.read_bytes(),
                                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
            assert response.status_code == 400
            assert "未找到价格列" in response.json()["detail"]
