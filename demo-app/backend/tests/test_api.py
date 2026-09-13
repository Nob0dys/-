from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.main import app


def workbook_bytes(unit: str = "台") -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "询价单"
    sheet.append(["序号", "产品名称", "参数", "型号", "单位", "数量"])
    sheet.append([1, "电子天平", "100g，0.001g，带防风罩", "", unit, 2])
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def login(client: TestClient):
    response = client.post("/api/auth/login", json={"username": "admin", "password": "admin123"})
    assert response.status_code == 200


def test_quote_job_persistence_candidates_and_exports():
    with TestClient(app) as client:
        assert client.get("/api/health").status_code == 200
        login(client)
        customers = client.get("/api/customers").json()
        ordinary = next(item for item in customers if item["customer_type"] == "ordinary")
        response = client.post(
            "/api/quote-jobs",
            data={"customer_id": ordinary["id"], "requested_option_count": 3},
            files={"file": ("询价.xlsx", workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert response.status_code == 202, response.text
        job_id = response.json()["id"]
        job = client.get(f"/api/quote-jobs/{job_id}").json()
        assert job["status"] == "review"
        assert job["total_lines"] == 1
        lines = client.get(f"/api/quote-jobs/{job_id}/lines?page=49&page_size=50").json()
        assert lines["page"] == 1
        line = client.get(f"/api/quote-lines/{lines['items'][0]['id']}").json()
        assert line["options"]
        selected = [item for item in line["options"] if item["selected"]]
        if not selected:
            selected = [line["options"][0]]
            update = client.patch(
                f"/api/quote-lines/{line['id']}",
                json={"selected_option_ids": [selected[0]["id"]], "final_prices": {}, "manual_note": "测试"},
            )
            assert update.status_code == 200
        confirm = client.post(f"/api/quote-lines/{line['id']}/confirm", json={"override_reason": "测试人工确认"})
        assert confirm.status_code == 200, confirm.text
        customer_export = client.post(f"/api/quote-jobs/{job_id}/export/customer")
        assert customer_export.status_code == 200
        workbook = load_workbook(BytesIO(customer_export.content))
        assert "多方案报价明细" in workbook.sheetnames
        joined = " ".join(str(cell.value or "") for row in workbook["多方案报价明细"] for cell in row)
        assert "历史来源" not in joined
        internal_export = client.post(f"/api/quote-jobs/{job_id}/export/internal")
        assert internal_export.status_code == 200
        internal = load_workbook(BytesIO(internal_export.content))
        assert "内部方案明细" in internal.sheetnames


def test_unit_conflict_filter_works_with_sqlite_json_encoding():
    with TestClient(app) as client:
        login(client)
        ordinary = next(
            item for item in client.get("/api/customers").json()
            if item["customer_type"] == "ordinary"
        )
        # 厂家规则：中文量词（个/支/台/套…）互变不阻断；物理量词（克/毫升…）
        # 与中文量词互变才构成阻断。询价单位用"克"（物理量词）触发冲突。
        response = client.post(
            "/api/quote-jobs",
            data={"customer_id": ordinary["id"], "requested_option_count": 3},
            files={"file": (
                "单位冲突.xlsx",
                workbook_bytes("克"),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )},
        )
        assert response.status_code == 202
        job_id = response.json()["id"]

        result = client.get(
            f"/api/quote-jobs/{job_id}/lines?row_filter=unit_conflict"
        ).json()
        assert result["total"] == 1
        assert "单位不可直接换算" in result["items"][0]["warnings"][0]


def test_vip_discount_requires_margin_approval_and_marks_preference():
    with TestClient(app) as client:
        login(client)
        customer_response = client.post(
            "/api/customers",
            json={
                "name": "VIP测试客户",
                "customer_type": "vip",
                "discount_percent": 8,
                "minimum_margin_percent": 15,
                "preferred_manufacturers": ["慈溪市华徐衡器实业有限公司"],
                "requirements": [],
            },
        )
        assert customer_response.status_code == 200
        customer = customer_response.json()
        job_response = client.post(
            "/api/quote-jobs",
            data={"customer_id": customer["id"], "requested_option_count": 3},
            files={"file": (
                "VIP询价.xlsx",
                workbook_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )},
        )
        job_id = job_response.json()["id"]
        result = client.get(f"/api/quote-jobs/{job_id}/lines").json()
        line = result["items"][0]

        assert line["status"] == "review"
        assert any("最低毛利线" in warning for warning in line["warnings"])
        detail = client.get(f"/api/quote-lines/{line['id']}").json()
        assert any(
            "VIP偏好制造商" in option["reasons"] for option in detail["options"]
        )


def test_special_customer_requires_structured_requirement():
    with TestClient(app) as client:
        login(client)
        response = client.post(
            "/api/customers",
            json={"name": "无条件特殊客户", "customer_type": "special", "requirements": []},
        )

        assert response.status_code == 422
