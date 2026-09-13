from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.main import app


def price_sheet_bytes(rows: list[list]) -> bytes:
    """构造专属报价单 xlsx：序号/名称/参数/单位/单价 表头。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "专属价目"
    sheet.append(["序号", "名称", "参数", "单位", "单价"])
    for row in rows:
        sheet.append(row)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def inquiry_bytes(name: str, spec: str, unit: str = "台") -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "询价单"
    sheet.append(["序号", "产品名称", "参数", "型号", "单位", "数量"])
    sheet.append([1, name, spec, "", unit, 1])
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def login(client: TestClient, username: str = "admin", password: str = "admin123"):
    response = client.post("/api/auth/login", json={"username": username, "password": password})
    assert response.status_code == 200, response.text


def create_customer(client: TestClient, name: str, **extra) -> int:
    response = client.post("/api/customers", json={"name": name, **extra})
    assert response.status_code == 200, response.text
    return response.json()["id"]


def upload_sheet(client: TestClient, customer_id: int, content: bytes, filename: str = "专属报价单.xlsx"):
    return client.post(
        f"/api/customers/{customer_id}/price-sheet",
        files={"file": (filename, content,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )


def create_job(client: TestClient, customer_id: int, content: bytes) -> str:
    response = client.post(
        "/api/quote-jobs",
        data={"customer_id": customer_id, "requested_option_count": 3},
        files={"file": ("询价.xlsx", content,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 202, response.text
    return response.json()["id"]


def first_line_detail(client: TestClient, job_id: str) -> dict:
    lines = client.get(f"/api/quote-jobs/{job_id}/lines").json()
    assert lines["total"] == 1, lines
    return client.get(f"/api/quote-lines/{lines['items'][0]['id']}").json()


def test_price_sheet_upload_get_replace_delete():
    """上传→查询→整表替换→删除 全流程；客户列表带 price_sheet_count。"""
    with TestClient(app) as client:
        login(client)
        customer_id = create_customer(client, "专属价目全流程客户")

        upload = upload_sheet(client, customer_id, price_sheet_bytes([
            [1, "专属协议天平", "量程200g，精度0.01g", "台", 66.0],
            [2, "专属协议天平", "量程200g，精度0.01g", "台", 66.0],  # 表内重复
            [3, "无价格产品", "缺单价", "台", None],  # 无效价格
            [4, "专属游标卡尺", "0-150mm", "把", 88.0],
        ]))
        assert upload.status_code == 200, upload.text
        assert upload.json() == {"inserted": 2, "skipped_duplicates": 1, "skipped_invalid": 1}

        sheet = client.get(f"/api/customers/{customer_id}/price-sheet").json()
        assert sheet["count"] == 2
        assert {row["name"] for row in sheet["rows"]} == {"专属协议天平", "专属游标卡尺"}
        assert all(set(row) >= {"name", "spec", "model", "brand", "manufacturer", "unit", "price"} for row in sheet["rows"])

        customers = client.get("/api/customers").json()
        assert next(item for item in customers if item["id"] == customer_id)["price_sheet_count"] == 2

        # 整表替换：旧行被删，只剩新行
        replaced = upload_sheet(client, customer_id, price_sheet_bytes([
            [1, "专属学生电源", "0-16V", "台", 120.0],
        ]))
        assert replaced.status_code == 200, replaced.text
        assert replaced.json()["inserted"] == 1
        sheet = client.get(f"/api/customers/{customer_id}/price-sheet").json()
        assert sheet["count"] == 1
        assert sheet["rows"][0]["name"] == "专属学生电源"

        deleted = client.delete(f"/api/customers/{customer_id}/price-sheet")
        assert deleted.status_code == 200, deleted.text
        assert deleted.json() == {"ok": True, "deleted": 1}
        sheet = client.get(f"/api/customers/{customer_id}/price-sheet").json()
        assert sheet == {"count": 0, "rows": []}
        # 重复删除为幂等空操作
        assert client.delete(f"/api/customers/{customer_id}/price-sheet").json() == {"ok": True, "deleted": 0}


def test_price_sheet_validation_and_auth():
    with TestClient(app) as client:
        login(client)
        customer_id = create_customer(client, "专属价目校验客户")

        bad_suffix = client.post(
            f"/api/customers/{customer_id}/price-sheet",
            files={"file": ("报价.txt", b"not excel", "text/plain")},
        )
        assert bad_suffix.status_code == 400

        no_price_column = BytesIO()
        workbook = Workbook()
        workbook.active.append(["序号", "名称", "参数"])
        workbook.active.append([1, "产品", "参数"])
        workbook.save(no_price_column)
        unparsable = upload_sheet(client, customer_id, no_price_column.getvalue())
        assert unparsable.status_code == 400

        assert upload_sheet(client, 999999, price_sheet_bytes([[1, "x", "y", "台", 1.0]])).status_code == 404
        assert client.get("/api/customers/999999/price-sheet").status_code == 404
        assert client.delete("/api/customers/999999/price-sheet").status_code == 404

    with TestClient(app) as quote_client:
        login(quote_client, "quote", "quote123")
        assert upload_sheet(quote_client, customer_id, price_sheet_bytes([[1, "x", "y", "台", 1.0]])).status_code == 403
        assert quote_client.get(f"/api/customers/{customer_id}/price-sheet").status_code == 403
        assert quote_client.delete(f"/api/customers/{customer_id}/price-sheet").status_code == 403


def test_exclusive_sheet_wins_default_selection_without_discount():
    """VIP 客户带协议折扣：专属价目候选优先默认选中，且 final_price 不再叠加折扣。"""
    with TestClient(app) as client:
        login(client)
        vip_id = create_customer(
            client, "专属价目VIP客户", customer_type="vip", discount_percent=10,
        )
        # 公共库同名竞品：100 元（验证专属行压过公共行，且公共行仍正常打折）
        public = client.post(
            "/api/history",
            json={"name": "专属协议天平", "spec": "量程200g，精度0.01g", "unit": "台",
                  "manufacturer": "公共厂", "price": 100.0},
        )
        assert public.status_code == 201, public.text

        upload = upload_sheet(client, vip_id, price_sheet_bytes([
            [1, "专属协议天平", "量程200g，精度0.01g", "台", 66.0],
        ]))
        assert upload.status_code == 200, upload.text
        assert upload.json()["inserted"] == 1

        job_id = create_job(client, vip_id, inquiry_bytes("专属协议天平", "量程200g，精度0.01g"))
        detail = first_line_detail(client, job_id)
        exclusive_options = [
            option for option in detail["options"]
            if "客户专属价目" in (option["record"].get("source_file") or "")
        ]
        assert exclusive_options, detail["options"]
        exclusive = exclusive_options[0]
        # 默认选中的是专属候选，价格为协议价 66 元（不打折成 59.4）
        assert exclusive["selected"] is True
        assert exclusive["final_price"] == 66.0
        assert "客户专属价目" in exclusive["reasons"]
        assert not any("已应用客户折扣" in str(w) for w in exclusive["warnings"])
        # 公共库候选仍作为备选展示，并正常叠加 10% 协议折扣
        public_options = [
            option for option in detail["options"]
            if option["record"].get("manufacturer") == "公共厂"
        ]
        assert public_options
        assert public_options[0]["final_price"] == 90.0
        assert any("已应用客户折扣" in str(w) for w in public_options[0]["warnings"])


def test_exclusive_rows_hidden_from_other_customers():
    """其它客户的报价任务候选池中绝不出现专属价目行。"""
    with TestClient(app) as client:
        login(client)
        owner_id = create_customer(client, "专属价目属主客户")
        upload = upload_sheet(client, owner_id, price_sheet_bytes([
            [1, "专属隔离量筒", "100ml", "个", 12.0],
        ]))
        assert upload.status_code == 200, upload.text

        other_id = create_customer(client, "无专属价目客户")
        job_id = create_job(client, other_id, inquiry_bytes("专属隔离量筒", "100ml", "个"))
        detail = first_line_detail(client, job_id)
        assert all(
            "客户专属价目" not in (option["record"].get("source_file") or "")
            for option in detail["options"]
        ), detail["options"]


def test_dedup_history_does_not_touch_exclusive_rows():
    """治理去重只作用于公共库：与公共行同 key 的专属行不被删，公共行也不被并掉。"""
    with TestClient(app) as client:
        login(client)
        owner_id = create_customer(client, "专属去重保护客户")
        upload = upload_sheet(client, owner_id, price_sheet_bytes([
            [1, "专属去重仪器", "规格D", "台", 50.0],
        ]))
        assert upload.status_code == 200, upload.text
        # 公共库造一条与专属行完全同 key（名称/规格/单位/制造商/价格）的记录
        public = client.post(
            "/api/history",
            json={"name": "专属去重仪器", "spec": "规格D", "unit": "台", "price": 50.0},
        )
        assert public.status_code == 201, public.text

        response = client.post("/api/governance/dedup-history")
        assert response.status_code == 200, response.text

        # 专属行仍在（count 不变），公共行也仍在——两者互不参与对方去重
        sheet = client.get(f"/api/customers/{owner_id}/price-sheet").json()
        assert sheet["count"] == 1
        found = client.get("/api/history/search", params={"q": "专属去重仪器"}).json()
        assert any(item["id"] == public.json()["id"] for item in found)


def test_price_draft_approve_not_blocked_by_exclusive_row():
    """草稿批准写入公共库时不与专属行去重：同 key 专属行存在也能正常写入。"""
    with TestClient(app) as client:
        login(client)
        owner_id = create_customer(client, "专属草稿回写客户")
        upload = upload_sheet(client, owner_id, price_sheet_bytes([
            [1, "草稿回写仪器", "规格E", "台", 30.0],
        ]))
        assert upload.status_code == 200, upload.text

        # 该客户的任务命中专属行并被默认选中；确认后产生人工确认价草稿，
        # 草稿去重 key 与专属行完全相同（旧逻辑会因此跳过公共库写入）
        job_id = create_job(client, owner_id, inquiry_bytes("草稿回写仪器", "规格E"))
        detail = first_line_detail(client, job_id)
        selected_ids = [option["id"] for option in detail["options"] if option["selected"]]
        assert selected_ids
        confirm = client.post(
            f"/api/quote-lines/{detail['id']}/confirm", json={"override_reason": ""}
        )
        assert confirm.status_code == 200, confirm.text

        drafts = client.get("/api/price-drafts", params={"status": "pending"}).json()
        draft = next(item for item in drafts if item["name"] == "草稿回写仪器")
        approved = client.post(f"/api/price-drafts/{draft['id']}/review", json={"action": "approve"})
        assert approved.status_code == 200, approved.text
        # 不被专属行阻塞：真正写入了公共库
        assert approved.json()["history_quote_id"] is not None
        found = client.get("/api/history/search", params={"q": "草稿回写仪器"}).json()
        assert any(item["id"] == approved.json()["history_quote_id"] for item in found)
        # 专属行不受影响
        assert client.get(f"/api/customers/{owner_id}/price-sheet").json()["count"] == 1
